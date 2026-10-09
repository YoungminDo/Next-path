"""Ingestion workbook (template v2.1) -> SEED through the ingestion pipeline. Synthetic workbook
with the template's exact headers; each test runs in a transaction that is rolled back."""
import copy
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import text

from hellomyme.ingest.sheet import import_workbook, write_report
from tests.conftest import AS_OF
from tests.test_ingest import _reference

H = {
    "01_PERSON": ["person_id", "member_id", "person_type", "career_stage", "country_scope",
                  "record_status", "created_at"],
    "02_EDUCATION": ["education_id", "person_id", "institution_raw", "institution_id", "major_raw",
                     "major_id", "degree_level", "graduation_year", "graduation_status", "source_id",
                     "review_status"],
    "03_WORK_EVENT": ["work_event_id", "person_id", "organization_raw", "organization_id", "role_raw",
                      "role_id", "industry_id", "employment_type", "start_value", "start_precision",
                      "end_value", "end_precision", "is_current", "is_parallel", "source_id",
                      "review_status"],
    "04_ORGANIZATION": ["organization_id", "canonical_name", "aliases", "industry_id", "company_type",
                        "size_band", "country", "review_status"],
    "05_INSTITUTION": ["institution_id", "canonical_name", "aliases", "institution_type", "country"],
    "06_ROLE_TAXONOMY": ["role_id", "parent_role_id", "role_family", "canonical_role", "level",
                         "aliases"],
    "07_INDUSTRY_TAXONOMY": ["industry_id", "parent_industry_id", "canonical_industry", "level",
                             "aliases"],
    "08_MAJOR_TAXONOMY": ["major_id", "parent_major_id", "canonical_major", "level", "aliases"],
    "09_SOURCE": ["source_id", "person_id", "source_type", "submission_id", "submitted_at",
                  "permitted_use", "storage_ref", "retention_until", "verification_level"],
    "12_CONSENT": ["consent_id", "person_id", "purpose", "consent_version", "granted",
                   "effective_at", "withdrawn_at"],
    "13_SNAPSHOT": ["snapshot_id", "person_id", "as_of_date", "career_months",
                    "current_work_event_id", "career_sequence", "education_snapshot",
                    "model_input_version"],
    "15_MENTOR": ["mentor_id", "person_id", "mentor_opt_in", "profile_visibility", "contact_enabled",
                  "available_topics"],
}

ROWS = {
    "01_PERSON": [
        ["P1", None, "member", "experienced", "KR", "active", "2026-10-08"],
        ["P2", None, "member", "student", "KR", "active", "2026-10-08"],
        ["P_DEMO", None, "member", "experienced", "KR", "demo", "2026-10-08"],
        ["P4", None, "member", "experienced", "KR", "active", "2026-10-08"],
    ],
    "02_EDUCATION": [
        ["E1", "P1", "인제스트대", "INST1", "인제스트학과", "M1", "bachelor", 2018, "graduated", "S1",
         "approved"],
        ["E2", "P2", "인제스트대학교", "INST1", "인제스트학과", "M1", "학사", 2027, "졸업예정", "S2",
         "pending"],
    ],
    "03_WORK_EVENT": [
        ["W1", "P1", "인제스트(주)", "ORG1", "그로스 마케터", "R1", None, "full_time", "2018-03", "month",
         datetime(2020, 2, 1), None, False, False, "S1", "approved"],
        ["W2", "P1", "시트신규", "ORG2", "브랜드 매니저", None, None, "intern", "2020-03", "month",
         "2020.08", "month", "FALSE", False, "S1", "검수완료"],
        ["W3", "P1", "미등록회사", None, "없는 직무", "R9", None, "founder", 2021, "year", None, None,
         "현재", False, "S1", "approved"],
        ["W4", "P1", "잘못 넣은 회사", None, "삭제할 직무", None, None, "employee", 2015, "year", 2016,
         "year", False, False, "S1", "rejected"],
        ["W5", "P2", "인제스트컴퍼니", "ORG1", "그로스 해커", None, None, "intern", "2025-07", "month",
         "2025-12", "month", False, False, "S2", "pending"],
    ],
    "04_ORGANIZATION": [
        ["ORG1", "인제스트컴퍼니", "인제스트(주), Ingest Co.", "I1", "corporation", "large", "KR",
         "approved"],
        ["ORG2", "시트신규회사", "시트신규", None, "startup", "small", "KR", "approved"],
    ],
    "05_INSTITUTION": [["INST1", "인제스트대학교", "인제스트대; Ingest Univ", "university", "KR"]],
    "06_ROLE_TAXONOMY": [["R1", None, "business", "인제스트 그로스", 2, "그로스 마케터"],
                         ["R9", None, "business", "없는 직무", 2, None]],
    "07_INDUSTRY_TAXONOMY": [["I1", None, "없는 산업", 1, None]],
    "08_MAJOR_TAXONOMY": [["M1", None, "인제스트학", 1, None]],
    "09_SOURCE": [
        ["S1", "P1", "linkedin_capture", "SUB1", "2026-10-01", "AGGREGATE_ONLY", "s3://cap/p1.png",
         "2027-10-01", "unverified"],
        ["S2", "P2", "resume", "SUB2", "2026-10-01", "AGGREGATE_ONLY", None, None, None],
        ["S4", "P4", "self_submitted", "SUB4", "2026-10-01", "AGGREGATE_ONLY", None, None, None],
    ],
    "12_CONSENT": [
        ["C2", "P2", "career_stats", "v1", True, "2026-10-01", None],
        ["C4", "P4", "career_stats", "v1", True, "2026-09-01", "2026-10-02"],
    ],
    "13_SNAPSHOT": [["SN1", "P1", "2026-10-01", 80, "W3", "W1>W2>W3", None, "v0"]],
    "15_MENTOR": [["MN1", "P1", True, "public", True, "마케팅"]],
}


def workbook(path: Path, rows=None, drop_column=None) -> Path:
    rows = rows or ROWS
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    guide = wb.create_sheet("00_GUIDE")
    guide.append(["항목", "설명", "규칙"])
    for title, header in H.items():
        ws = wb.create_sheet(title)
        cols = [c for c in header if c != drop_column]
        ws.append(cols)
        for r in rows.get(title, []):
            ws.append([v for c, v in zip(header, r, strict=True) if c != drop_column])
    wb.save(path)
    return path


@pytest.fixture
def db(engine):
    with engine.connect() as c:
        tx = c.begin()
        _reference(c)
        yield c
        tx.rollback()


def run(db, path, **kw):
    return import_workbook(db, path, as_of=AS_OF, legal_basis="test basis", **kw)


def people(report):
    return {e["person_id"]: e for e in report["people"]}


def codes(entry):
    return {x["code"] for x in entry["findings"]}


def test_workbook_loads_reviewed_people_and_holds_the_rest(db, tmp_path):
    report = run(db, workbook(tmp_path / "wb.xlsx"))
    by = people(report)
    assert {k: v["status"] for k, v in by.items()} == {
        "P1": "LOADED", "P2": "NEEDS_REVIEW", "P_DEMO": "SKIPPED", "P4": "WITHDRAWN"}
    assert (by["P1"]["work_events"], by["P1"]["educations"]) == (3, 1)  # W4 was rejected
    assert {"MENTOR_IGNORED"} <= codes(by["P1"])
    assert {"HELD_FOR_REVIEW", "GRADUATION_NOT_FINAL"} <= codes(by["P2"])
    assert "DERIVED_SHEET_IGNORED" in {x["code"] for x in report["findings"]}
    assert "INDUSTRY_UNMATCHED" in {x["code"] for x in report["findings"]}

    jobs = db.execute(text(
        """SELECT w.organization_raw, o.name AS org, n.code AS role, w.event_type, w.employment_type,
                  w.start_date::text, w.start_date_precision, w.end_date::text, w.end_date_precision,
                  w.is_current, w.normalization_status
           FROM work_event w LEFT JOIN organization o USING (organization_id)
           LEFT JOIN taxonomy_node n ON n.taxonomy_node_id = w.role_taxonomy_node_id
           WHERE w.person_id = :p ORDER BY w.start_date"""),
        {"p": by["P1"]["loaded_person_id"]}).all()
    assert [tuple(j) for j in jobs] == [
        # raw name kept, canonical from 04_ORGANIZATION; role through the 06 alias; an Excel date cell
        ("인제스트(주)", "인제스트컴퍼니", "ING_R2", "EMPLOYMENT", "FULL_TIME", "2018-03-01", "MONTH",
         "2020-02-01", "MONTH", False, "MAPPED"),
        # a new organisation from the team's list; the unknown role waits in the queue
        ("시트신규", "시트신규회사", None, "EMPLOYMENT", "INTERN", "2020-03-01", "MONTH", "2020-08-01",
         "MONTH", False, "PARTIAL"),
        ("미등록회사", None, None, "STARTUP", None, "2021-01-01", "YEAR", None, None, True, "UNMAPPED"),
    ]
    edu = db.execute(text(
        """SELECT e.institution_raw, i.name, n.code, e.degree_type, e.graduation_year,
                  p.declared_stage, p.origin_layer
           FROM education e JOIN institution i USING (institution_id)
           JOIN taxonomy_node n ON n.taxonomy_node_id = e.major_taxonomy_node_id
           JOIN person p USING (person_id) WHERE e.person_id = :p"""),
        {"p": by["P1"]["loaded_person_id"]}).one()
    assert tuple(edu) == ("인제스트대", "인제스트대학교", "ING_M1", "BACHELOR", 2018, "PROFESSIONAL", "SEED")
    sub = db.execute(text(
        """SELECT s.source_type, s.permitted_use, s.legal_basis, a.storage_ref
           FROM source_submission s JOIN source_asset a USING (source_submission_id)
           WHERE s.external_submission_id = 'P1'""")).one()
    assert tuple(sub) == ("PUBLIC_PROFILE", "AGGREGATE_ONLY", "test basis", "s3://cap/p1.png")
    assert db.execute(text(
        "SELECT legal_basis FROM source_submission WHERE external_submission_id = 'P2'")).scalar() == (
        "CONSENT:career_stats:v1")
    # Nothing at all is stored for the withdrawn or the demo person.
    assert not db.execute(text(
        "SELECT 1 FROM source_submission WHERE external_submission_id IN ('P4', 'P_DEMO')")).first()

    # Aliases from the reference sheets now standardise future input on their own.
    from hellomyme.ingest.pipeline import _lookup
    assert _lookup(db, "ORGANIZATION", "Ingest Co.") == _lookup(db, "ORGANIZATION", "인제스트컴퍼니")
    assert _lookup(db, "INSTITUTION", "Ingest Univ") is not None
    queued = set(db.execute(text(
        "SELECT entity_kind, raw_value FROM mapping_queue WHERE status = 'PENDING'")).all())
    assert {("ROLE", "없는 직무"), ("ORGANIZATION", "미등록회사"), ("ROLE", "브랜드 매니저")} <= queued


def test_reupload_is_idempotent_and_picks_up_new_reviews(db, tmp_path):
    first = run(db, workbook(tmp_path / "a.xlsx"))
    again = run(db, workbook(tmp_path / "b.xlsx"))
    assert people(again)["P1"]["status"] == "ALREADY_LOADED"
    assert people(again)["P1"]["loaded_person_id"] == people(first)["P1"]["loaded_person_id"]
    assert people(again)["P2"]["status"] == "NEEDS_REVIEW"
    count = "SELECT count(*) FROM person WHERE origin_layer = 'SEED'"
    before = db.execute(text(count)).scalar_one()

    rows = copy.deepcopy(ROWS)  # the team finishes reviewing P2
    for sheet in ("02_EDUCATION", "03_WORK_EVENT"):
        for r in rows[sheet]:
            if r[1] == "P2":
                r[-1] = "approved"
    third = run(db, workbook(tmp_path / "c.xlsx", rows))
    assert people(third)["P2"]["status"] == "LOADED"
    assert db.execute(text(count)).scalar_one() == before + 1
    stage, grad = db.execute(text(
        """SELECT p.declared_stage, e.graduation_year FROM person p JOIN education e USING (person_id)
           WHERE p.person_id = :p"""), {"p": people(third)["P2"]["loaded_person_id"]}).one()
    assert (stage, grad) == ("STUDENT", None)  # an expected graduation is not a graduation


def test_accept_unreviewed_loads_pending_rows(db, tmp_path):
    report = run(db, workbook(tmp_path / "wb.xlsx"), accept_unreviewed=True)
    assert people(report)["P2"]["status"] == "LOADED"


@pytest.mark.parametrize("sheet, idx, col, value, code", [
    ("09_SOURCE", 0, 2, "카톡으로 받음", "UNKNOWN_SOURCE_TYPE"),
    ("09_SOURCE", 0, 5, None, "MISSING_PERMITTED_USE"),
    ("03_WORK_EVENT", 0, 14, "S_NOPE", "UNKNOWN_SOURCE"),
    ("03_WORK_EVENT", 0, 14, "S2", "SOURCE_PERSON_MISMATCH"),
    ("03_WORK_EVENT", 0, 8, "2018년 13월", "BAD_DATE"),
    ("03_WORK_EVENT", 2, 9, "month", "PRECISION_MISMATCH"),  # 2021 has no month to keep
])
def test_sheet_errors_block_only_that_person(db, tmp_path, sheet, idx, col, value, code):
    rows = copy.deepcopy(ROWS)
    rows[sheet][idx][col] = value
    report = run(db, workbook(tmp_path / "wb.xlsx", rows))
    p1 = people(report)["P1"]
    assert p1["status"] == "BLOCKED"
    hit = next(x for x in p1["findings"] if x["code"] == code)
    assert hit["severity"] == "BLOCK" and hit["sheet"] == sheet
    assert hit["row"] == idx + 2  # the row the team sees in Excel
    assert people(report)["P2"]["status"] == "NEEDS_REVIEW"  # others are unaffected


def test_record_rule_rejects_after_conversion(db, tmp_path):
    rows = copy.deepcopy(ROWS)
    rows["03_WORK_EVENT"][0][10] = "2017-01"  # ends before it starts
    p1 = people(run(db, workbook(tmp_path / "wb.xlsx", rows)))["P1"]
    assert p1["status"] == "REJECTED"
    hit = next(x for x in p1["findings"] if x["code"] == "END_BEFORE_START")
    assert (hit["sheet"], hit["row"], hit["ref"]) == ("03_WORK_EVENT", 2, "W1")


def test_unknown_employment_type_waits_for_a_person(db, tmp_path):
    rows = copy.deepcopy(ROWS)
    rows["03_WORK_EVENT"][0][7] = None
    p1 = people(run(db, workbook(tmp_path / "wb.xlsx", rows)))["P1"]
    assert p1["status"] == "NEEDS_REVIEW" and "EVENT_TYPE_UNKNOWN" in codes(p1)


def test_missing_column_stops_the_workbook(db, tmp_path):
    report = run(db, workbook(tmp_path / "wb.xlsx", drop_column="permitted_use"))
    assert report["people"] == []
    assert report["findings"][0]["code"] == "MISSING_COLUMN"
    assert not db.execute(text("SELECT 1 FROM source_submission")).first()


def test_report_workbook_lists_what_to_fix(db, tmp_path):
    rows = copy.deepcopy(ROWS)
    rows["03_WORK_EVENT"][0][8] = "언젠가"
    report = run(db, workbook(tmp_path / "wb.xlsx", rows))
    out = write_report(report, tmp_path / "report.xlsx", dry_run=True)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["요약", "사람별 결과", "고칠 것", "표준화 대기"]
    fixes = list(wb["고칠 것"].iter_rows(min_row=2, values_only=True))
    assert fixes[0][:3] == ("막힘", "03_WORK_EVENT", 2)
    results = {r[0]: r[2] for r in wb["사람별 결과"].iter_rows(min_row=2, values_only=True)}
    assert results == {"P1": "거부(시트 오류)", "P2": "검수 대기", "P_DEMO": "건너뜀", "P4": "동의 철회"}
