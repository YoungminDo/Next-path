"""PRE_SEED v1.3 workbook (HELLOMYME_PreSeed_Career_Data_v1.3.xlsx) -> Package.

The workbook is the official format from v1.3 on. Its sheets are mapped onto the same staged
row shape the CSV importer validates, plus the v1.4 extras (taxonomy nodes, gender, admission
year, date precision, role/major/industry nodes). Values are only re-typed to text; nothing is
filled in or corrected here.
"""
from __future__ import annotations

import hashlib
from datetime import date, datetime
from pathlib import Path

import openpyxl

from hellomyme.importer.preseed import Package

SHEETS = {
    "readme": "00_README", "persons": "01_PERSON", "educations": "02_EDUCATION",
    "work_events": "03_WORK_EVENT", "institutions": "04_INSTITUTION",
    "organizations": "05_ORGANIZATION", "taxonomies": "06_TAXONOMY",
    "taxonomy_nodes": "07_TAXONOMY_NODE", "org_industries": "08_ORG_INDUSTRY",
    "work_event_sources": "09_WORK_EVENT_SOURCE",
}


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sheet(wb, title: str) -> list[dict]:
    if title not in wb.sheetnames:
        raise FileNotFoundError(f"missing sheet {title}")
    rows = wb[title].iter_rows(values_only=True)
    header = [_text(h) for h in next(rows)]
    return [dict(zip(header, (_text(v) for v in r), strict=False))
            for r in rows if any(v is not None for v in r)]


def readme(path: str | Path) -> dict[str, str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return {r.get("Item", ""): r.get("Value", "") for r in _sheet(wb, SHEETS["readme"])}
    finally:
        wb.close()


def load_workbook_package(path: str | Path) -> Package:
    path = Path(path)
    data = path.read_bytes()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        raw = {name: _sheet(wb, title) for name, title in SHEETS.items() if name != "readme"}
    finally:
        wb.close()

    nodes = {n["taxonomy_node_id"]: n for n in raw["taxonomy_nodes"]}

    def family(node_id: str) -> str:
        """Display name of the depth-2 ancestor (legacy major_family / job_family)."""
        node = nodes.get(node_id)
        while node and node.get("depth") not in ("", "1", "2") and node.get("parent_node_id"):
            node = nodes.get(node["parent_node_id"])
        return (node or {}).get("display_name", "")

    # Legacy major/role dictionaries are keyed by the raw text, which is what the source said.
    majors: dict[str, dict] = {}
    educations = []
    for e in raw["educations"]:
        name = e.get("raw_major_name", "")
        if name and name not in majors:
            majors[name] = {"major_id": name, "major_name": name,
                            "major_family": family(e.get("major_taxonomy_node_id", ""))}
        educations.append({**e, "major_id": name})

    roles: dict[str, dict] = {}
    work_events = []
    for w in raw["work_events"]:
        title = w.get("raw_role_title", "")
        if title and title not in roles:
            roles[title] = {"role_id": title, "role_name": title,
                            "job_family": family(w.get("role_taxonomy_node_id", ""))}
        work_events.append({**w, "role_id": title})

    rows = {
        "institutions": raw["institutions"], "majors": list(majors.values()),
        "roles": list(roles.values()), "organizations": raw["organizations"],
        "persons": raw["persons"], "educations": educations, "work_events": work_events,
        "work_event_sources": raw["work_event_sources"],
        "taxonomies": raw["taxonomies"], "taxonomy_nodes": raw["taxonomy_nodes"],
        "org_industries": raw["org_industries"],
    }
    # Fixed key: renaming the file must not make the same workbook look like a new package.
    return Package(rows=rows, manifest={"workbook.xlsx": hashlib.sha256(data).hexdigest()})
