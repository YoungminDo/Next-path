"""Ingestion workbook (HELLOMYME_Career_Ingestion_Template v2.1) -> pipeline -> SEED.

The workbook is the input standard: people type and review career data in it, this module loads
it. Each person (01_PERSON row) becomes one submission through the same pipeline as AI
extractions (pipeline.receive / approve), so the same rules hold: raw names kept, unknown names
queued for standardisation, rows not yet reviewed held back, re-imports idempotent.

Sheets used: 01 PERSON, 02 EDUCATION, 03 WORK_EVENT, 04 ORGANIZATION, 05 INSTITUTION,
06 ROLE / 07 INDUSTRY / 08 MAJOR taxonomy, 09 SOURCE, 12 CONSENT, 15 MENTOR (opt-in is ignored:
SEED people are never shown as mentors). 13/14/16 are computed outputs and never read as input;
10/11/17/18 are working logs and are not applied. See docs/09_INGESTION_PIPELINE.md §10.
"""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
from sqlalchemy import Connection, text

from hellomyme.ingest import pipeline
from hellomyme.ingest.pipeline import IngestError, Issue

TEMPLATE = "career_ingestion_v2.1"
SHEETS = {"PERSON", "EDUCATION", "WORK_EVENT", "ORGANIZATION", "INSTITUTION", "ROLE_TAXONOMY",
          "INDUSTRY_TAXONOMY", "MAJOR_TAXONOMY", "SOURCE", "CONSENT", "MENTOR"}
REQUIRED = {"PERSON": ("person_id",), "EDUCATION": ("person_id", "source_id"),
            "WORK_EVENT": ("person_id", "source_id"),
            "SOURCE": ("source_id", "person_id", "source_type", "permitted_use")}
DERIVED = {"SNAPSHOT", "TRANSITION", "SIMILARITY"}
NOT_APPLIED = {"PARSE_JOB", "FIELD_REVIEW", "MAPPING_QUEUE", "QA_LOG"}


def _k(v) -> str:
    return re.sub(r"[\s_\-./()]+", "", str(v).strip().lower()) if v is not None else ""


def _vocab(table: dict[str, tuple[str, ...]]) -> dict[str, str]:
    return {_k(word): code for code, words in table.items() for word in words}


# Values the workbook may use, in Korean or English, in any case or spacing.
EMPLOYMENT = _vocab({
    "EMPLOYMENT": ("employee", "employed", "직원", "재직", "회사원", "employment"),
    "EMPLOYMENT/FULL_TIME": ("full_time", "fulltime", "full-time", "정규직", "permanent"),
    "EMPLOYMENT/INTERN": ("intern", "internship", "인턴", "인턴십", "체험형인턴", "채용연계형인턴"),
    "EMPLOYMENT/CONTRACT": ("contract", "contractor", "계약직", "파견", "temporary"),
    "EMPLOYMENT/PART_TIME": ("part_time", "parttime", "파트타임", "아르바이트", "알바"),
    "FREELANCE": ("freelance", "freelancer", "프리랜서"),
    "STARTUP": ("founder", "cofounder", "co-founder", "startup", "창업", "창업자", "공동창업"),
    "SELF_EMPLOYED": ("self_employed", "selfemployed", "자영업", "개인사업"),
    "SIDE_BUSINESS": ("side_business", "sideproject", "부업", "사이드프로젝트"),
    "MILITARY": ("military", "군복무", "군대", "병역"),
    "CAREER_BREAK": ("career_break", "break", "휴직", "공백", "경력단절"),
    "STUDY": ("study", "학업", "대학원", "어학연수", "유학"),
    "PROJECT": ("project", "프로젝트"),
    "OTHER": ("other", "기타"),
})
DEGREE = _vocab({
    "ASSOCIATE": ("associate", "전문학사", "전문대"), "BACHELOR": ("bachelor", "학사", "학부", "ba", "bs"),
    "MASTER": ("master", "석사", "mba", "ms", "ma"), "DOCTORATE": ("doctorate", "phd", "박사"),
    "OTHER": ("other", "기타"),
})
GRADUATION = _vocab({
    "FINAL": ("graduated", "graduate", "졸업", "completed", "수료"),
    "NOT_FINAL": ("expected", "졸업예정", "enrolled", "재학", "attending", "휴학", "leave", "dropped",
                  "dropout", "중퇴", "자퇴", "withdrawn", "제적"),
})
STAGE = _vocab({
    "PROFESSIONAL": ("experienced", "professional", "worker", "employed", "경력", "경력자", "직장인"),
    "STUDENT": ("student", "재학", "재학생", "대학생"),
    "JOB_SEEKER": ("job_seeker", "jobseeker", "new_grad", "newgrad", "취준", "취업준비", "신입"),
})
SOURCE_TYPE = _vocab({
    "PUBLIC_PROFILE": ("public_profile", "linkedin", "linkedin_capture", "linkedin_screenshot",
                       "링크드인", "링크드인캡처", "공개프로필"),
    "LINKEDIN_USER_UPLOAD": ("linkedin_upload", "linkedin_export", "user_upload", "self_upload",
                             "본인업로드", "링크드인업로드"),
    "CAREER_DOCUMENT": ("career_document", "resume", "cv", "document", "이력서", "경력기술서",
                        "경력증명서", "self_submitted", "direct", "direct_submission", "직접제출",
                        "동의제출", "consent"),
    "MANUAL_RESEARCH": ("manual_research", "manual", "research", "리서치", "수기"),
})
# When one person has several sources, the strictest one decides how the person may be used.
SOURCE_PRECEDENCE = ["PUBLIC_PROFILE", "MANUAL_RESEARCH", "LINKEDIN_USER_UPLOAD", "CAREER_DOCUMENT"]
REVIEW = _vocab({
    "OK": ("approved", "approve", "reviewed", "verified", "done", "ok", "confirmed", "accepted",
           "검수완료", "승인", "완료", "확인", "확정"),
    "SKIP": ("rejected", "reject", "excluded", "exclude", "deleted", "반려", "삭제", "제외"),
    "HOLD": ("pending", "review", "in_review", "needs_review", "todo", "검수전", "대기", "검토중"),
})
PERSON_SKIP = {_k(w) for w in ("demo", "test", "sample", "예시", "테스트", "excluded", "exclude",
                               "withdrawn", "deleted", "inactive", "삭제", "제외")}
SIZE = _vocab({"MICRO": ("micro", "초소형"), "SMALL": ("small", "소기업", "스타트업"),
               "MID": ("mid", "medium", "중견", "중기업"), "LARGE": ("large", "대기업"),
               "ENTERPRISE": ("enterprise", "그룹", "대기업집단")})
TRUE = {_k(w) for w in ("true", "t", "y", "yes", "1", "o", "예", "네", "현재", "재직중", "current")}
FALSE = {_k(w) for w in ("false", "f", "n", "no", "0", "x", "아니오", "아니요")}

MESSAGES = {
    "SCHEMA": "형식이 맞지 않아요",
    "DUPLICATE_LOCAL_ID": "같은 ID가 두 번 있어요",
    "DUPLICATE_ID": "같은 ID가 두 번 있어요",
    "EVIDENCE_ASSET_UNKNOWN": "source_id가 09_SOURCE에 없어요",
    "PRECISION_MISMATCH": "정밀도(year/month)와 날짜 값이 맞지 않아요",
    "YEAR_OUT_OF_RANGE": "연도가 1950~2100 밖이에요",
    "END_BEFORE_START": "종료가 시작보다 빨라요",
    "START_AFTER_AS_OF": "시작이 기준일보다 미래예요",
    "CURRENT_WITH_END_DATE": "현재 재직(is_current)인데 종료일이 있어요",
    "MISSING_START": "시작 시점이 비어 있어요",
    "EVENT_TYPE_UNKNOWN": "employment_type을 알 수 없어요 (예: employee, intern, contract, founder)",
    "DUPLICATE_EVENT": "같은 경력(회사·직무·시작)이 반복돼 하나만 남겼어요",
    "BAD_DATE": "날짜를 읽을 수 없어요 (YYYY 또는 YYYY-MM)",
    "BAD_BOOLEAN": "참/거짓 값을 읽을 수 없어 비워뒀어요 (TRUE/FALSE)",
    "UNKNOWN_VALUE": "알 수 없는 값이라 비워뒀어요",
    "UNKNOWN_SOURCE": "09_SOURCE에 없는 source_id예요",
    "SOURCE_PERSON_MISMATCH": "source_id가 다른 사람의 출처예요",
    "UNKNOWN_SOURCE_TYPE": "source_type을 알 수 없어요 (예: linkedin_capture, resume, self_submitted)",
    "NO_SOURCE": "이 사람의 출처(09_SOURCE)가 없어요",
    "MISSING_PERMITTED_USE": "permitted_use가 비어 있어요",
    "CONFLICTING_PERMITTED_USE": "한 사람의 출처들에 permitted_use가 서로 달라요",
    "NO_LEGAL_BASIS": "동의(12_CONSENT)도 --legal-basis도 없어요",
    "CONSENT_WITHDRAWN": "동의가 철회돼 적재하지 않았어요",
    "WITHDRAWN_AFTER_LOAD": "이미 적재된 사람의 동의가 철회됐어요: 삭제 절차가 필요해요",
    "UNKNOWN_PERSON": "01_PERSON에 없는 person_id예요",
    "UNKNOWN_REFERENCE": "참조한 ID가 기준 시트에 없어 이름으로 찾았어요",
    "NO_INSTITUTION": "학교 이름과 institution_id가 모두 비어 있어요",
    "NO_ORG_OR_ROLE": "회사와 직무가 모두 비어 있어요",
    "UNKNOWN_REVIEW_STATUS": "review_status를 알 수 없어 검수 대기로 뒀어요",
    "HELD_FOR_REVIEW": "review_status가 검수 전이라 대기 중이에요",
    "GRADUATION_NOT_FINAL": "졸업이 확정되지 않아 졸업연도는 넣지 않았어요",
    "MEMBER_ID_IGNORED": "member_id는 쓰지 않아요 (회원 연결은 본인 동의 후)",
    "MENTOR_IGNORED": "멘토는 가입 후 본인 opt-in으로만 열려요 (SEED는 멘토로 노출하지 않음)",
    "DERIVED_SHEET_IGNORED": "계산 결과 시트라 입력으로 쓰지 않아요",
    "SHEET_NOT_APPLIED": "작업 기록 시트라 적재에 반영하지 않아요",
    "MISSING_SHEET": "필요한 시트가 없어요",
    "MISSING_COLUMN": "필요한 열이 없어요",
    "AMBIGUOUS_NAME": "같은 이름이 여러 항목과 맞아 연결하지 않았어요 (표준화 대기)",
    "TAXONOMY_UNMATCHED": "분류표에 없는 항목이라 표준화 대기열에 넣었어요",
    "INDUSTRY_UNMATCHED": "산업 분류표에 없는 항목이라 연결하지 않았어요",
    "LOADED_CHANGED": "이미 적재된 사람인데 내용이 바뀌었어요: 수정 적재는 다음 단계예요",
    "DAY_PRECISION_REDUCED": "일 단위는 월 단위로 저장해요",
}


@dataclass
class Row:
    sheet: str
    row: int
    values: dict

    def get(self, col: str):
        return self.values.get(col)


@dataclass
class Finding:
    severity: str  # BLOCK | WARN | INFO
    code: str
    sheet: str | None = None
    row: int | None = None
    ref: str | None = None
    person_id: str | None = None
    detail: str | None = None

    @property
    def message(self) -> str:
        return MESSAGES.get(self.code, self.code) + (f" ({self.detail})" if self.detail else "")


@dataclass
class Workbook:
    path: Path
    sha256: str
    sheets: dict[str, list[Row]]
    titles: dict[str, str]
    findings: list[Finding] = field(default_factory=list)


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def read_workbook(path: str | Path) -> Workbook:
    path = Path(path)
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, list[Row]] = {}
    titles: dict[str, str] = {}
    findings: list[Finding] = []
    for ws in wb.worksheets:
        name = re.sub(r"^\d+[_\s.-]*", "", ws.title).strip().upper().replace(" ", "_")
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip().lower() if h is not None else None for h in rows[0]]
        data = []
        for i, r in enumerate(rows[1:], start=2):
            values = {h: _clean(v) for h, v in zip(header, r, strict=False) if h}
            if any(v is not None for v in values.values()):
                data.append(Row(ws.title, i, values))
        if name in DERIVED and data:
            findings.append(Finding("INFO", "DERIVED_SHEET_IGNORED", ws.title))
        elif name in NOT_APPLIED and data:
            findings.append(Finding("INFO", "SHEET_NOT_APPLIED", ws.title))
        if name in SHEETS:
            sheets[name], titles[name] = data, ws.title
            for col in REQUIRED.get(name, ()):
                if col not in header:
                    findings.append(Finding("BLOCK", "MISSING_COLUMN", ws.title, 1, detail=col))
    wb.close()
    for need in ("PERSON", "WORK_EVENT", "SOURCE"):
        if need not in sheets:
            findings.append(Finding("BLOCK", "MISSING_SHEET", detail=need))
    return Workbook(path, sha, sheets, titles, findings)


# --- value readers --------------------------------------------------------------------------

def _partial(value, precision, f: list[Finding], r: Row, col: str) -> dict | None:
    """'2018', 2018, '2018-03', '2018.3', '2018년 3월' or an Excel date -> partial date, with the
    precision the workbook states. A month is never invented."""
    if value is None:
        return None
    prec = {"year": "YEAR", "y": "YEAR", "연": "YEAR", "년": "YEAR", "month": "MONTH", "m": "MONTH",
            "월": "MONTH", "day": "DAY", "d": "DAY", "일": "DAY"}.get(_k(precision)) if precision else None
    y = m = None
    if isinstance(value, (datetime, date)):
        y, m = value.year, value.month
        prec = prec or "MONTH"
    elif isinstance(value, (int, float)) and float(value).is_integer():
        y = int(value)
    else:
        hit = re.fullmatch(r"\s*(\d{4})\s*(?:[-./년]\s*(\d{1,2})\s*월?(?:[-./]\s*\d{1,2}\s*일?)?)?\s*",
                           str(value))
        if not hit or (hit.group(2) and not 1 <= int(hit.group(2)) <= 12):
            f.append(Finding("BLOCK", "BAD_DATE", r.sheet, r.row, detail=f"{col}={value}"))
            return None
        y, m = int(hit.group(1)), int(hit.group(2)) if hit.group(2) else None
    if prec == "DAY":
        f.append(Finding("INFO", "DAY_PRECISION_REDUCED", r.sheet, r.row, detail=col))
        prec = "MONTH"
    prec = prec or ("MONTH" if m else "YEAR")
    if prec == "MONTH" and m is None:
        f.append(Finding("BLOCK", "PRECISION_MISMATCH", r.sheet, r.row, detail=col))
        return None
    return {"value": f"{y:04d}-{m:02d}" if prec == "MONTH" else f"{y:04d}", "precision": prec}


def _bool(value, f: list[Finding], r: Row, col: str) -> bool | None:
    if value is None or isinstance(value, bool):
        return value
    key = _k(value)
    if key in TRUE:
        return True
    if key in FALSE:
        return False
    f.append(Finding("WARN", "BAD_BOOLEAN", r.sheet, r.row, detail=f"{col}={value}"))
    return None


def _code(vocab: dict, value, f: list[Finding], r: Row, col: str) -> str | None:
    if value is None:
        return None
    code = vocab.get(_k(value))
    if code is None:
        f.append(Finding("WARN", "UNKNOWN_VALUE", r.sheet, r.row, detail=f"{col}={value}"))
    return code


def _aliases(value) -> list[str]:
    return [a.strip() for a in re.split(r"[,;|\n]", str(value or "")) if a.strip()]


def _jsonable(values: dict) -> dict:
    return {k: v.isoformat() if isinstance(v, (date, datetime)) else v for k, v in values.items()}


# --- reference sheets -------------------------------------------------------------------------

def _version(conn: Connection, taxonomy: str):
    return conn.execute(text(
        """INSERT INTO taxonomy_version (taxonomy, version, description)
           VALUES (:t, 'ingest', 'Created by the career ingestion pipeline')
           ON CONFLICT (taxonomy, version) DO UPDATE SET description = EXCLUDED.description
           RETURNING taxonomy_version_id"""), {"t": taxonomy}).scalar_one()


def _references(conn: Connection, wb: Workbook, source_id: str) -> dict[str, dict[str, str]]:
    """Workbook reference ids (ORG_001, R_012, ...) -> database ids. Organisations and
    institutions are matched by name or alias and created when new; taxonomy entries are only
    matched to the active product taxonomy (never created) and queued when unknown."""
    f = wb.findings
    out: dict[str, dict[str, str]] = {k: {} for k in ("ORGANIZATION", "INSTITUTION", "ROLE",
                                                       "INDUSTRY", "MAJOR")}

    def unique(rows, id_col):
        seen = {}
        for r in rows:
            rid = r.get(id_col)
            if rid is None:
                continue
            if rid in seen:
                f.append(Finding("BLOCK", "DUPLICATE_ID", r.sheet, r.row, str(rid)))
            seen[rid] = r
        return seen

    def resolve(kind: str, names: list[str], r: Row, rid) -> tuple[str | None, bool]:
        """(id, ambiguous): the first name with any match decides; a tie is never resolved."""
        for n in names:
            found = pipeline._candidates(conn, kind, n)
            if len(found) == 1:
                return next(iter(found)), False
            if found:
                f.append(Finding("WARN", "AMBIGUOUS_NAME", r.sheet, r.row, str(rid), detail=n))
                pipeline._queue(conn, kind, n) if kind != "INDUSTRY" else None
                return None, True
        return None, False

    for kind, sheet, id_col, name_col in (("ROLE", "ROLE_TAXONOMY", "role_id", "canonical_role"),
                                          ("INDUSTRY", "INDUSTRY_TAXONOMY", "industry_id",
                                           "canonical_industry"),
                                          ("MAJOR", "MAJOR_TAXONOMY", "major_id", "canonical_major")):
        for rid, r in unique(wb.sheets.get(sheet, []), id_col).items():
            names = [str(n) for n in [r.get(name_col), *_aliases(r.get("aliases"))] if n]
            node, ambiguous = resolve(kind, names, r, rid)
            if node is None:
                if ambiguous:
                    continue
                if kind == "INDUSTRY":
                    f.append(Finding("WARN", "INDUSTRY_UNMATCHED", r.sheet, r.row, str(rid),
                                     detail=r.get(name_col)))
                elif r.get(name_col):
                    pipeline._queue(conn, kind, str(r.get(name_col)))
                    f.append(Finding("INFO", "TAXONOMY_UNMATCHED", r.sheet, r.row, str(rid),
                                     detail=r.get(name_col)))
                continue
            out[kind][rid] = node
            for alias in names:
                conn.execute(text(
                    """INSERT INTO taxonomy_alias (taxonomy_node_id, alias_text, source_type)
                       VALUES (CAST(:n AS uuid), :a, 'INGEST_WORKBOOK') ON CONFLICT DO NOTHING"""),
                    {"n": node, "a": alias})

    for kind, sheet, id_col in (("ORGANIZATION", "ORGANIZATION", "organization_id"),
                                ("INSTITUTION", "INSTITUTION", "institution_id")):
        table = kind.lower()
        for rid, r in unique(wb.sheets.get(sheet, []), id_col).items():
            if REVIEW.get(_k(r.get("review_status"))) == "SKIP" or not r.get("canonical_name"):
                continue
            names = [str(n) for n in [r.get("canonical_name"), *_aliases(r.get("aliases"))]]
            found, ambiguous = resolve(kind, names, r, rid)
            if ambiguous:
                continue
            if found is None:  # a new entry from the team's curated list
                cols = {"name": names[0], "taxonomy_version_id": _version(conn, kind)}
                if kind == "ORGANIZATION":
                    cols["company_size_band"] = _code(SIZE, r.get("size_band"), f, r, "size_band")
                else:
                    cols["institution_type"] = _str(r.get("institution_type"))
                found = conn.execute(text(
                    f"""INSERT INTO {table} ({", ".join(cols)})
                        VALUES ({", ".join(f":{c}" for c in cols)})
                        ON CONFLICT (name) DO UPDATE SET name = EXCLUDED.name
                        RETURNING {table}_id::text"""), cols).scalar_one()
            out[kind][rid] = found
            for alias in names:
                conn.execute(text(
                    f"""INSERT INTO {table}_alias ({table}_id, alias_text, source_type)
                        VALUES (CAST(:i AS uuid), :a, 'INGEST_WORKBOOK') ON CONFLICT DO NOTHING"""),
                    {"i": found, "a": alias})
            node = out["INDUSTRY"].get(r.get("industry_id")) if kind == "ORGANIZATION" else None
            if node:
                conn.execute(text(
                    """INSERT INTO organization_industry (organization_id, taxonomy_node_id,
                           is_primary, source_id)
                       SELECT CAST(:o AS uuid), CAST(:n AS uuid),
                              NOT EXISTS (SELECT 1 FROM organization_industry
                                          WHERE organization_id = CAST(:o AS uuid) AND is_primary
                                            AND valid_to IS NULL), CAST(:s AS uuid)
                       WHERE NOT EXISTS (SELECT 1 FROM organization_industry
                                         WHERE organization_id = CAST(:o AS uuid)
                                           AND taxonomy_node_id = CAST(:n AS uuid))"""),
                    {"o": found, "n": node, "s": source_id})
    return out


# --- people -----------------------------------------------------------------------------------

@dataclass
class PersonPlan:
    person_id: str
    row: Row
    doc: dict | None = None
    refs: dict = field(default_factory=dict)
    codes: dict = field(default_factory=dict)
    held: set = field(default_factory=set)
    rows: dict = field(default_factory=dict)  # local_id -> Row
    source_type: str | None = None
    permitted_use: str | None = None
    legal_basis: str | None = None
    withdrawn: bool = False
    findings: list[Finding] = field(default_factory=list)


def _plan_person(p: PersonPlan, wb: Workbook, refs: dict, by_person: dict, *,
                 legal_basis: str | None, accept_unreviewed: bool) -> None:
    f = p.findings
    pr = p.row
    if pr.get("member_id"):
        f.append(Finding("WARN", "MEMBER_ID_IGNORED", pr.sheet, pr.row, p.person_id))
    stage = _code(STAGE, pr.get("career_stage"), f, pr, "career_stage")

    # Provenance: every row must point at a source of this person.
    sources = {r.get("source_id"): r for r in by_person["SOURCE"].get(p.person_id, [])}
    all_sources = {r.get("source_id"): r for r in wb.sheets.get("SOURCE", [])}
    if not sources:
        f.append(Finding("BLOCK", "NO_SOURCE", pr.sheet, pr.row, p.person_id))
    types, uses = [], set()
    for sid, r in sources.items():
        t = SOURCE_TYPE.get(_k(r.get("source_type")))
        if t is None:
            f.append(Finding("BLOCK", "UNKNOWN_SOURCE_TYPE", r.sheet, r.row, str(sid),
                             detail=str(r.get("source_type"))))
        else:
            types.append(t)
        if r.get("permitted_use"):
            uses.add(str(r.get("permitted_use")))
        else:
            f.append(Finding("BLOCK", "MISSING_PERMITTED_USE", r.sheet, r.row, str(sid)))
    if len(uses) > 1:
        f.append(Finding("BLOCK", "CONFLICTING_PERMITTED_USE", pr.sheet, pr.row, p.person_id,
                         detail=", ".join(sorted(uses))))
    p.source_type = min(types, key=SOURCE_PRECEDENCE.index) if types else None
    p.permitted_use = next(iter(uses)) if len(uses) == 1 else None

    # Consent: a withdrawal (or a refusal) wins over everything; otherwise consent is the basis.
    consents = by_person["CONSENT"].get(p.person_id, [])
    granted = []
    for r in consents:
        ok = _bool(r.get("granted"), f, r, "granted")
        if r.get("withdrawn_at") or ok is False:
            p.withdrawn = True
        elif ok:
            granted.append(r)
    if granted:
        r = granted[-1]
        p.legal_basis = "CONSENT:" + ":".join(str(x) for x in (r.get("purpose"), r.get("consent_version"))
                                               if x is not None)
    else:
        p.legal_basis = legal_basis
    if p.legal_basis is None and not p.withdrawn:
        f.append(Finding("BLOCK", "NO_LEGAL_BASIS", pr.sheet, pr.row, p.person_id))
    for r in by_person["MENTOR"].get(p.person_id, []):
        if _bool(r.get("mentor_opt_in"), f, r, "mentor_opt_in"):
            f.append(Finding("INFO", "MENTOR_IGNORED", r.sheet, r.row, p.person_id))

    def evidence(r: Row, rid: str):
        sid = r.get("source_id")
        if sid is None or sid not in all_sources:
            f.append(Finding("BLOCK", "UNKNOWN_SOURCE", r.sheet, r.row, rid, detail=str(sid)))
        elif sid not in sources:
            f.append(Finding("BLOCK", "SOURCE_PERSON_MISMATCH", r.sheet, r.row, rid, detail=str(sid)))
        return [{"asset_id": str(sid or "?"), "text": f"{r.sheet} {r.row}행 ({rid})"}]

    def status(r: Row, entity: str, rid: str) -> str:
        raw = r.get("review_status")
        s = REVIEW.get(_k(raw)) if raw is not None else "HOLD"
        if s is None:
            f.append(Finding("WARN", "UNKNOWN_REVIEW_STATUS", r.sheet, r.row, rid, detail=str(raw)))
            s = "HOLD"
        if s == "HOLD" and accept_unreviewed:
            s = "OK"
        if s == "HOLD":
            p.held.add((entity, rid))
            f.append(Finding("INFO", "HELD_FOR_REVIEW", r.sheet, r.row, rid))
        return s

    def ref(kind: str, rid: str, value, r: Row, col: str):
        if value is None:
            return None, None
        hit = refs[kind].get(value)
        if hit is None and _canonical(wb, kind, value) is None:
            f.append(Finding("WARN", "UNKNOWN_REFERENCE", r.sheet, r.row, rid, detail=f"{col}={value}"))
        return hit, value

    educations = []
    for r in by_person["EDUCATION"].get(p.person_id, []):
        rid = str(r.get("education_id") or f"E@{r.row}")
        if status(r, "EDUCATION", rid) == "SKIP":
            continue
        inst_ref, inst_key = ref("INSTITUTION", rid, r.get("institution_id"), r, "institution_id")
        major_ref, _ = ref("MAJOR", rid, r.get("major_id"), r, "major_id")
        inst_raw = r.get("institution_raw") or _canonical(wb, "INSTITUTION", inst_key)
        if not inst_raw:
            f.append(Finding("BLOCK", "NO_INSTITUTION", r.sheet, r.row, rid))
            continue
        grad = r.get("graduation_year")
        g = GRADUATION.get(_k(r.get("graduation_status"))) if r.get("graduation_status") else "FINAL"
        if g == "NOT_FINAL" and grad is not None:
            f.append(Finding("INFO", "GRADUATION_NOT_FINAL", r.sheet, r.row, rid,
                             detail=str(r.get("graduation_status"))))
            grad = None
        elif g is None:
            _code(GRADUATION, r.get("graduation_status"), f, r, "graduation_status")
        e = {"local_id": rid, "institution_raw": str(inst_raw),
             "major_raw": _str(r.get("major_raw")) or _canonical(wb, "MAJOR", r.get("major_id")),
             "degree_raw": _str(r.get("degree_level")),
             "start": _partial(r.get("admission_year"), "year", f, r, "admission_year"),
             "end": _partial(grad, "year", f, r, "graduation_year"),
             "evidence": evidence(r, rid)}
        educations.append(e)
        p.rows[rid] = r
        if inst_ref:
            p.refs[("EDUCATION", rid, "institution")] = inst_ref
        if major_ref:
            p.refs[("EDUCATION", rid, "major")] = major_ref
        p.codes[("EDUCATION", rid, "degree_type")] = _code(DEGREE, r.get("degree_level"), f, r,
                                                            "degree_level")

    events = []
    for r in by_person["WORK_EVENT"].get(p.person_id, []):
        rid = str(r.get("work_event_id") or f"W@{r.row}")
        if status(r, "WORK_EVENT", rid) == "SKIP":
            continue
        org_ref, org_key = ref("ORGANIZATION", rid, r.get("organization_id"), r, "organization_id")
        role_ref, role_key = ref("ROLE", rid, r.get("role_id"), r, "role_id")
        org = _str(r.get("organization_raw")) or _canonical(wb, "ORGANIZATION", org_key)
        role = _str(r.get("role_raw")) or _canonical(wb, "ROLE", role_key)
        if not org and not role:
            f.append(Finding("BLOCK", "NO_ORG_OR_ROLE", r.sheet, r.row, rid))
            continue
        kind = r.get("event_type") or r.get("employment_type")
        typed = EMPLOYMENT.get(_k(kind)) if kind is not None else None
        if kind is not None and typed is None:
            f.append(Finding("WARN", "UNKNOWN_VALUE", r.sheet, r.row, rid, detail=f"employment_type={kind}"))
        event_type, _, emp = (typed or "").partition("/")
        if r.get("event_type") and r.get("employment_type"):  # both given: type + detail
            detail = EMPLOYMENT.get(_k(r.get("employment_type")), "")
            emp = detail.partition("/")[2] or emp
        w = {"local_id": rid, "organization_raw": org, "role_raw": role,
             "employment_type_raw": _str(r.get("employment_type")),
             "event_type_hint": event_type or "UNKNOWN",
             "start": _partial(r.get("start_value"), r.get("start_precision"), f, r, "start_value"),
             "end": _partial(r.get("end_value"), r.get("end_precision"), f, r, "end_value"),
             "is_current": _bool(r.get("is_current"), f, r, "is_current"),
             "evidence": evidence(r, rid)}
        events.append({k: v for k, v in w.items() if v is not None or k in ("start", "end")})
        p.rows[rid] = r
        if org_ref:
            p.refs[("WORK_EVENT", rid, "organization")] = org_ref
        if role_ref:
            p.refs[("WORK_EVENT", rid, "role")] = role_ref
        p.codes[("WORK_EVENT", rid, "employment_type")] = emp or None

    used = {e["evidence"][0]["asset_id"] for e in educations + events}
    assets = [{"asset_id": str(sid), "page_order": i,
               **({"storage_ref": str(r.get("storage_ref"))} if r.get("storage_ref") else {})}
              for i, (sid, r) in enumerate(sorted(sources.items(), key=lambda x: x[1].row), start=1)]
    assets += [{"asset_id": a, "page_order": len(assets) + i}
               for i, a in enumerate(sorted(used - {a["asset_id"] for a in assets}), start=1)]
    p.doc = {
        "schema_version": "career_extraction.v1",
        "submission_id": p.person_id,
        "assets": assets or [{"asset_id": "WORKBOOK", "page_order": 1}],
        "person": {"declared_stage": stage} if stage else {},
        "educations": educations,
        "work_events": events,
        "sheet": {"template": TEMPLATE, "accept_unreviewed": accept_unreviewed,
                  "rows": [{"sheet": r.sheet, "row": r.row, "values": _jsonable(r.values)}
                           for r in [pr, *sources.values(), *p.rows.values()]]},
    }


def _str(v):
    return None if v is None else str(v)


def _canonical(wb: Workbook, kind: str, ref_id) -> str | None:
    if ref_id is None:
        return None
    sheet, id_col, name_col = {
        "ORGANIZATION": ("ORGANIZATION", "organization_id", "canonical_name"),
        "INSTITUTION": ("INSTITUTION", "institution_id", "canonical_name"),
        "ROLE": ("ROLE_TAXONOMY", "role_id", "canonical_role"),
        "MAJOR": ("MAJOR_TAXONOMY", "major_id", "canonical_major")}[kind]
    hit = next((r for r in wb.sheets.get(sheet, []) if r.get(id_col) == ref_id), None)
    return _str(hit.get(name_col)) if hit else None


# --- run --------------------------------------------------------------------------------------

def import_workbook(conn: Connection, path: str | Path, *, as_of: date, collector: str = "workbook",
                    legal_basis: str | None = None, accept_unreviewed: bool = False) -> dict:
    wb = read_workbook(path)
    report = {"workbook": Path(path).name, "sha256": wb.sha256, "template": TEMPLATE,
              "as_of": as_of.isoformat(), "people": [], "findings": []}
    if any(x.severity == "BLOCK" for x in wb.findings):
        report["findings"] = [asdict(x) | {"message": x.message} for x in wb.findings]
        report["summary"] = {"people": 0, "blocked": "workbook"}
        return report
    # The reference sheets themselves are kept verbatim as one append-only source record.
    raw = {name: [{"row": r.row, "values": _jsonable(r.values)} for r in wb.sheets.get(name, [])]
           for name in ("ORGANIZATION", "INSTITUTION", "ROLE_TAXONOMY", "INDUSTRY_TAXONOMY",
                        "MAJOR_TAXONOMY")}
    params = {"k": wb.sha256, "raw": json.dumps(raw, ensure_ascii=False),
              "meta": json.dumps({"template": TEMPLATE, "workbook": Path(path).name})}
    source_id = conn.execute(text(
        """INSERT INTO source_record (source_type, data_layer, source_system, source_key,
               raw_payload, metadata)
           VALUES ('MANUAL_RESEARCH', 'SEED', 'HELLOMYME_INGEST_WORKBOOK', :k, CAST(:raw AS jsonb),
                   CAST(:meta AS jsonb))
           ON CONFLICT (source_system, source_key) DO NOTHING RETURNING source_id::text"""),
        params).scalar() or conn.execute(text(
        """SELECT source_id::text FROM source_record
           WHERE source_system = 'HELLOMYME_INGEST_WORKBOOK' AND source_key = :k"""), params).scalar_one()
    refs = _references(conn, wb, source_id)

    by_person: dict[str, dict[str, list[Row]]] = {k: {} for k in ("SOURCE", "CONSENT", "MENTOR",
                                                                   "EDUCATION", "WORK_EVENT")}
    people_ids = {r.get("person_id") for r in wb.sheets.get("PERSON", [])}
    for sheet in by_person:
        for r in wb.sheets.get(sheet, []):
            pid = r.get("person_id")
            if pid not in people_ids:
                wb.findings.append(Finding("WARN", "UNKNOWN_PERSON", r.sheet, r.row, str(pid)))
                continue
            by_person[sheet].setdefault(pid, []).append(r)

    seen = set()
    for pr in wb.sheets.get("PERSON", []):
        pid = pr.get("person_id")
        if pid is None:
            continue
        pid = str(pid)
        entry = {"person_id": pid, "sheet_row": pr.row, "status": None, "parse_run_id": None,
                 "loaded_person_id": None, "findings": []}
        report["people"].append(entry)
        if pid in seen:
            entry["status"] = "SKIPPED"
            entry["findings"].append(Finding("BLOCK", "DUPLICATE_ID", pr.sheet, pr.row, pid))
            continue
        seen.add(pid)
        if _k(pr.get("record_status")) in PERSON_SKIP:
            entry["status"] = "SKIPPED"
            entry["reason"] = f"record_status={pr.get('record_status')}"
            continue
        p = PersonPlan(pid, pr)
        _plan_person(p, wb, refs, by_person, legal_basis=legal_basis,
                     accept_unreviewed=accept_unreviewed)
        entry["findings"] = p.findings
        if p.withdrawn:
            loaded = conn.execute(text(
                """SELECT s.loaded_person_id::text
                   FROM source_submission s JOIN collector c USING (collector_id)
                   WHERE c.name = :c AND s.external_submission_id = :p AND s.status = 'LOADED'"""),
                {"c": collector, "p": pid}).scalar()
            entry["status"] = "WITHDRAWN"
            p.findings.append(Finding("WARN" if loaded else "INFO",
                                      "WITHDRAWN_AFTER_LOAD" if loaded else "CONSENT_WITHDRAWN",
                                      pr.sheet, pr.row, pid))
            entry["loaded_person_id"] = loaded
            continue
        if any(x.severity == "BLOCK" for x in p.findings):
            entry["status"] = "BLOCKED"
            continue
        extra = [Issue(x.severity, x.code, None, x.ref, {"sheet": x.sheet, "row": x.row})
                 for x in p.findings if x.severity in ("WARN", "INFO")]
        try:
            out = pipeline.receive(
                conn, json.dumps(p.doc, ensure_ascii=False), collector=collector,
                source_type=p.source_type, permitted_use=p.permitted_use,
                legal_basis=p.legal_basis, model_version=f"workbook:{TEMPLATE}",
                prompt_version=f"sha256:{wb.sha256[:16]}", as_of=as_of, trusted=True,
                refs=p.refs, held=p.held, codes=p.codes, extra_issues=extra)
        except IngestError as exc:
            entry["status"] = "LOADED_CHANGED" if "already LOADED" in str(exc) else "ERROR"
            p.findings.append(Finding("WARN", "LOADED_CHANGED", pr.sheet, pr.row, pid, detail=str(exc))
                              if entry["status"] == "LOADED_CHANGED"
                              else Finding("BLOCK", "SCHEMA", pr.sheet, pr.row, pid, detail=str(exc)))
            continue
        entry["parse_run_id"] = out["parse_run_id"]
        for i in out["issues"]:
            if "sheet" in i["detail"]:
                continue  # one of the workbook findings above, passed through
            r = p.rows.get(i["local_id"])
            p.findings.append(Finding(i["severity"], i["code"], r.sheet if r else pr.sheet,
                                      r.row if r else pr.row, i["local_id"],
                                      detail=i["detail"].get("message") or i["detail"].get("field")))
        if out["status"] == "LOADED":
            entry["status"] = "ALREADY_LOADED"
            entry["loaded_person_id"] = conn.execute(text(
                """SELECT loaded_person_id::text FROM source_submission s JOIN parse_run r
                       USING (source_submission_id) WHERE r.parse_run_id = CAST(:r AS uuid)"""),
                {"r": out["parse_run_id"]}).scalar()
        elif out["status"] == "READY":
            loaded = pipeline.approve(conn, out["parse_run_id"])
            entry["status"] = "LOADED"
            entry["loaded_person_id"] = loaded["person_id"]
            entry["work_events"] = len(loaded["work_event_ids"])
            entry["educations"] = len(loaded["education_ids"])
        else:
            entry["status"] = out["status"]  # NEEDS_REVIEW or REJECTED
            entry["pending_fields"] = out.get("pending")

    for e in report["people"]:
        e["findings"] = [asdict(x) | {"message": x.message} for x in e["findings"]]
    report["findings"] = [asdict(x) | {"message": x.message} for x in wb.findings]
    counts: dict[str, int] = {}
    for e in report["people"]:
        counts[e["status"]] = counts.get(e["status"], 0) + 1
    report["summary"] = {"people": len(report["people"]), **counts}
    report["mapping_queue"] = [dict(r) for r in conn.execute(text(
        """SELECT entity_kind, raw_value, occurrences FROM mapping_queue WHERE status = 'PENDING'
           ORDER BY occurrences DESC, entity_kind, raw_value LIMIT 500""")).mappings()]
    return report


STATUS_KO = {"people": "사람 수", "LOADED": "적재됨", "ALREADY_LOADED": "이미 적재됨(변경 없음)",
             "NEEDS_REVIEW": "검수 대기",
             "REJECTED": "거부(규칙 위반)", "BLOCKED": "거부(시트 오류)", "SKIPPED": "건너뜀",
             "WITHDRAWN": "동의 철회", "LOADED_CHANGED": "적재 후 변경됨", "ERROR": "오류"}


def write_report(report: dict, path: str | Path, *, dry_run: bool) -> Path:
    """The same report as an xlsx the operations team can work from: what loaded, what is
    waiting and why (sheet + row), and which names need standardising."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.append(["항목", "값"])
    ws.append(["파일", report["workbook"]])
    ws.append(["모드", "미리보기(저장 안 함)" if dry_run else "적재"])
    ws.append(["기준일", report["as_of"]])
    for k, v in report.get("summary", {}).items():
        ws.append([STATUS_KO.get(k, k), v])
    ws = wb.create_sheet("사람별 결과")
    ws.append(["person_id", "행", "결과", "경력 수", "학력 수", "검수 대기 필드", "적재 ID", "비고"])
    for e in report["people"]:
        ws.append([e["person_id"], e["sheet_row"], STATUS_KO.get(e["status"], e["status"]),
                   e.get("work_events"), e.get("educations"), e.get("pending_fields"),
                   e.get("loaded_person_id"), e.get("reason")])
    ws = wb.create_sheet("고칠 것")
    ws.append(["심각도", "시트", "행", "ID", "person_id", "코드", "설명"])
    sev = {"BLOCK": "막힘", "WARN": "확인", "INFO": "참고"}
    rows = [(x, None) for x in report["findings"]] + [
        (x, e["person_id"]) for e in report["people"] for x in e["findings"]]
    order = {"BLOCK": 0, "WARN": 1, "INFO": 2}
    for x, pid in sorted(rows, key=lambda t: (order[t[0]["severity"]], t[0]["sheet"] or "",
                                              t[0]["row"] or 0)):
        ws.append([sev[x["severity"]], x["sheet"], x["row"], x["ref"], pid or x.get("person_id"),
                   x["code"], x["message"]])
    ws = wb.create_sheet("표준화 대기")
    ws.append(["종류", "이름(원문)", "등장 횟수"])
    for q in report.get("mapping_queue", []):
        ws.append([q["entity_kind"], q["raw_value"], q["occurrences"]])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    path = Path(path)
    wb.save(path)
    return path
